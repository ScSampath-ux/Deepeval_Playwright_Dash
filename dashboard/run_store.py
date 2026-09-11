"""
ShipConsole AI Evaluation Command Center - Run Persistence & Insights Engine

Transforms a raw DeepEval run (.deepeval/.latest_run_full.json) into a persisted,
navigable "Evaluation Run" record under reports/runs/run_<timestamp>.json, and
computes rule-based (non-LLM) insights and run-over-run deltas.

This module has no dependency on Flask or any dashboard web code, so it can be
imported directly from the evaluation pipeline (evals/test_agent_synthesized.py)
without pulling in the web server. Deleting the dashboard/ folder entirely only
removes the Command Center UI - it does not affect evaluation execution, since
the pipeline guards this import in a try/except.
"""

import os
import re
import json
import glob

# ---------------------------------------------------------------------------
# Metric Category Mapping (static - mirrors evals/metrics/suite.py groupings)
# ---------------------------------------------------------------------------
METRIC_CATEGORIES = {
    "Bias": "safety",
    "Toxicity": "safety",
    "PII Leakage": "safety",
    "Answer Relevancy": "quality",
    "Task Completion": "quality",
    "Contextual Recall": "quality",
    "Faithfulness": "quality",
    "Contextual Precision": "quality",
    "Contextual Relevancy": "quality",
    "Business Rule Compliance": "business",
    "HazMat & Dangerous Goods Compliance": "business",
    "Consolidation & Freight Compliance": "business",
    "Bot Response Verification (No Fallback Errors)": "business",
}

CATEGORY_LABELS = {
    "safety": "Safety & Guardrails",
    "quality": "Answer Quality",
    "business": "Business Rules & Compliance",
    "other": "Other",
}


def _clean_metric_name(name: str) -> str:
    """Strips the '[GEval]' suffix DeepEval appends to custom GEval metric names."""
    if not name:
        return name
    return re.sub(r"\s*\[GEval\]\s*$", "", name.strip())


def _categorize(name: str) -> str:
    return METRIC_CATEGORIES.get(_clean_metric_name(name), "other")


def runs_dir(reports_dir: str) -> str:
    d = os.path.join(reports_dir, "runs")
    os.makedirs(d, exist_ok=True)
    return d


def list_run_files(reports_dir: str):
    """Returns run_*.json paths under reports/runs, sorted oldest -> newest by filename."""
    d = runs_dir(reports_dir)
    files = glob.glob(os.path.join(d, "run_*.json"))
    files.sort()
    return files


def load_run(path: str):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def _build_metrics_summary(test_cases):
    """Aggregates per-metric stats (avg score, threshold, pass/fail counts, category)."""
    stats = {}
    order = []
    for tc in test_cases:
        for m in tc.get("metrics", []):
            raw_name = m["name"]
            if raw_name not in stats:
                stats[raw_name] = {
                    "name": raw_name,
                    "category": _categorize(raw_name),
                    "threshold": m.get("threshold", 0.0),
                    "scores": [],
                    "passed": 0,
                    "failed": 0,
                }
                order.append(raw_name)
            entry = stats[raw_name]
            score = m.get("score")
            if score is not None:
                entry["scores"].append(float(score))
            if m.get("success", True):
                entry["passed"] += 1
            else:
                entry["failed"] += 1

    summary = []
    for name in order:
        s = stats[name]
        avg = round(sum(s["scores"]) / len(s["scores"]), 4) if s["scores"] else 0.0
        total = s["passed"] + s["failed"]
        pass_rate = round((s["passed"] / total * 100), 1) if total else 0.0
        summary.append({
            "name": name,
            "category": s["category"],
            "avg_score": avg,
            "threshold": s["threshold"],
            "passed": s["passed"],
            "failed": s["failed"],
            "pass_rate": pass_rate,
        })
    return summary


def build_run_record(latest_run_path: str, run_id: str, timestamp_iso: str,
                      model_used: str = None, dataset_source: str = None,
                      trace_urls_by_input: dict = None):
    """
    Reads DeepEval's raw .latest_run_full.json and transforms it into the
    persisted Run schema described in the PRD (section 3.2).

    Returns None if there is no usable data (e.g. run produced zero test cases).
    """
    if not os.path.exists(latest_run_path) or os.path.getsize(latest_run_path) == 0:
        return None

    with open(latest_run_path, "r", encoding="utf-8") as f:
        raw = json.load(f)

    raw_cases = raw.get("testCases") or raw.get("testCasesData") or []
    if not raw_cases:
        return None

    trace_urls_by_input = trace_urls_by_input or {}

    test_cases = []
    for idx, tc in enumerate(raw_cases, 1):
        raw_metrics = tc.get("metricsData") or tc.get("metrics") or []
        metrics = []
        for m in raw_metrics:
            score = m.get("score")
            threshold = m.get("threshold", 0.0)
            metrics.append({
                "name": _clean_metric_name(m.get("name", "Metric")),
                "category": _categorize(m.get("name", "")),
                "score": round(float(score), 4) if score is not None else None,
                "threshold": threshold,
                "success": bool(m.get("success", True)),
                "reason": m.get("reason") or "",
                "evaluation_model": m.get("evaluationModel"),
            })

        case_success = all(m["success"] for m in metrics) if metrics else bool(tc.get("success", True))
        input_text = tc.get("input", "")
        test_cases.append({
            "id": f"TC_{idx:03d}",
            "input": input_text,
            "expected_output": tc.get("expectedOutput") or tc.get("expected_output"),
            "actual_output": tc.get("actualOutput") or tc.get("actual_output"),
            "success": case_success,
            "duration_seconds": round(float(tc.get("runDuration")), 2) if tc.get("runDuration") is not None else None,
            "metrics": metrics,
            "langfuse_trace_url": trace_urls_by_input.get(input_text.strip()),
        })

    total_cases = len(test_cases)
    passed_cases = sum(1 for tc in test_cases if tc["success"])
    failed_cases = total_cases - passed_cases
    pass_rate = round((passed_cases / total_cases * 100), 1) if total_cases else 0.0

    record = {
        "id": run_id,
        "timestamp": timestamp_iso,
        "model_used": model_used or "unknown",
        "dataset_source": dataset_source or "unknown",
        "total_cases": total_cases,
        "passed_cases": passed_cases,
        "failed_cases": failed_cases,
        "overall_pass_rate": pass_rate,
        "metrics_summary": _build_metrics_summary(test_cases),
        "test_cases": test_cases,
    }
    return record


def persist_run(latest_run_path: str, reports_dir: str, run_id: str, timestamp_iso: str,
                 model_used: str = None, dataset_source: str = None, trace_urls_by_input: dict = None):
    """
    Builds a Run record from the raw DeepEval output and writes it to
    reports/runs/run_<run_id>.json. Returns the written record, or None if
    there was no data to persist.
    """
    record = build_run_record(latest_run_path, run_id, timestamp_iso, model_used, dataset_source,
                               trace_urls_by_input=trace_urls_by_input)
    if record is None:
        return None

    out_dir = runs_dir(reports_dir)
    out_path = os.path.join(out_dir, f"run_{run_id}.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(record, f, indent=2, ensure_ascii=False)

    print(f"[Dashboard] Persisted evaluation run to: {out_path}")
    return record


# ---------------------------------------------------------------------------
# Rule-Based Insights (no LLM calls - pure comparison against the prior run)
# ---------------------------------------------------------------------------

def compute_insights(current: dict, previous: dict = None):
    """
    Computes procedural insights for a run: ship/no-ship decision, failing
    metrics/cases, and regressions vs the previous run. No LLM calls.
    """
    pass_rate = current["overall_pass_rate"]
    if pass_rate >= 80:
        status = "healthy"
        status_label = "Healthy (Ready to Ship)"
    elif pass_rate >= 50:
        status = "warning"
        status_label = "Needs Attention"
    else:
        status = "critical"
        status_label = "Critical (Block Release)"

    can_ship = pass_rate >= 80

    failing_metrics = [m["name"] for m in current["metrics_summary"] if m["failed"] > 0]
    failing_case_ids = [tc["id"] for tc in current["test_cases"] if not tc["success"]]

    # Ranked by pass_rate (not raw avg_score): avg_score direction is metric-specific
    # (e.g. Bias/Toxicity are "lower is better"), but pass_rate already normalizes
    # for that via each metric's own success/threshold semantics.
    lowest_metrics = sorted(
        (m for m in current["metrics_summary"] if m["pass_rate"] < 100),
        key=lambda m: m["pass_rate"],
    )[:3]

    regressions = []
    improvements = []
    if previous:
        prev_by_name = {m["name"]: m for m in previous.get("metrics_summary", [])}
        for m in current["metrics_summary"]:
            prev = prev_by_name.get(m["name"])
            if not prev:
                continue
            delta = round(m["pass_rate"] - prev["pass_rate"], 1)
            if delta < 0:
                regressions.append({"name": m["name"], "delta": delta,
                                     "previous_pass_rate": prev["pass_rate"], "current_pass_rate": m["pass_rate"]})
            elif delta > 0:
                improvements.append({"name": m["name"], "delta": delta,
                                      "previous_pass_rate": prev["pass_rate"], "current_pass_rate": m["pass_rate"]})
        regressions.sort(key=lambda r: r["delta"])
        improvements.sort(key=lambda r: -r["delta"])

    pass_rate_delta = round(pass_rate - previous["overall_pass_rate"], 1) if previous else None

    return {
        "status": status,
        "status_label": status_label,
        "can_ship": can_ship,
        "pass_rate": pass_rate,
        "pass_rate_delta": pass_rate_delta,
        "failing_metrics": failing_metrics,
        "failing_case_ids": failing_case_ids,
        "lowest_metrics": [{"name": m["name"], "avg_score": m["avg_score"], "pass_rate": m["pass_rate"]} for m in lowest_metrics],
        "regressions": regressions,
        "improvements": improvements,
    }


def export_case_to_excel(record: dict, case_id: str):
    """
    Builds a single-case Excel workbook (input/expected/actual + per-metric
    breakdown) for the given case in the given run record. Returns raw bytes,
    or None if the case_id doesn't exist in this run.
    """
    case = next((tc for tc in record["test_cases"] if tc["id"] == case_id), None)
    if case is None:
        return None

    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    pass_font = Font(name="Calibri", size=10, bold=True, color="375623")
    fail_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    fail_font = Font(name="Calibri", size=10, bold=True, color="C65911")
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(*(Side(style="thin", color="D9D9D9") for _ in range(4)))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = case_id

    ws.cell(row=1, column=1, value=f"ShipConsole AI - Test Case {case_id} ({record['id']})").font = Font(
        size=14, bold=True, color="1F4E78")

    fields = [
        ("Input Prompt", case["input"]),
        ("Expected Output", case["expected_output"] or ""),
        ("Actual Output", case["actual_output"] or ""),
        ("Overall Status", "PASS" if case["success"] else "FAIL"),
    ]
    r = 3
    for label, value in fields:
        c1 = ws.cell(row=r, column=1, value=label)
        c2 = ws.cell(row=r, column=2, value=value)
        c1.font = Font(bold=True)
        c1.alignment = align_left
        c2.alignment = align_left
        c1.border = thin_border
        c2.border = thin_border
        if label == "Overall Status":
            c2.fill = pass_fill if case["success"] else fail_fill
            c2.font = pass_font if case["success"] else fail_font
        r += 1

    r += 1
    headers = ["Metric", "Score", "Threshold", "Result", "Reason"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
    r += 1

    for m in case["metrics"]:
        c1 = ws.cell(row=r, column=1, value=m["name"])
        c2 = ws.cell(row=r, column=2, value=m["score"])
        c3 = ws.cell(row=r, column=3, value=m["threshold"])
        c4 = ws.cell(row=r, column=4, value="PASS" if m["success"] else "FAIL")
        c5 = ws.cell(row=r, column=5, value=m["reason"])
        for c in (c1, c2, c3, c4, c5):
            c.border = thin_border
            c.alignment = align_left
        c2.alignment = align_center
        c3.alignment = align_center
        c4.alignment = align_center
        c4.fill = pass_fill if m["success"] else fail_fill
        c4.font = pass_font if m["success"] else fail_font
        r += 1

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 3, 12), 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def get_run_summary(record: dict):
    """Lightweight summary shape used for run list cards / history charts."""
    return {
        "id": record["id"],
        "timestamp": record["timestamp"],
        "model_used": record.get("model_used"),
        "dataset_source": record.get("dataset_source"),
        "total_cases": record["total_cases"],
        "passed_cases": record["passed_cases"],
        "failed_cases": record["failed_cases"],
        "overall_pass_rate": record["overall_pass_rate"],
    }
